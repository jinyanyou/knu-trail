"""'3학년 이상 개설 전공과목 성적 평가방법 변경 지정 내역'(학사과 공지 엑셀) → data/rules/relative_grading/<학기>.json

지정된 과목은 '상대평가 권장 과목'으로, 수강생의 50% 이내에서 A등급을 줄 수 있다 (학업성적처리규정 제3조 2항).
학기마다 새로 공지되므로 학기별 파일로 쌓아 둔다 — 새 공지가 나오면 엑셀만 넣고 다시 실행:

    python scripts/fetch_timetable.py 2027 1                      # 그 학기 개설 강좌 (분반 대조용, 먼저)
    python scripts/import_relative_grading.py <엑셀 파일>          # 학기·작성 기준일은 엑셀 제목에서 읽음
    python scripts/import_relative_grading.py <엑셀> 2027-1 2027-01-15   # 못 읽으면 직접

적용 단위는 ① 교과목번호 ② 개설대학 ③ 개설학과가 모두 같은 분반 전체 — 같은 과목이라도 다른 학과가 연 분반은 아니다.
그래서 같은 학기 개설 강좌(data/timetable/<학기>.json)에서 세 가지가 맞는 분반을 찾아 그 개설학과 코드도 함께 적는다
(개설 강좌 파일이 없으면 학과 이름으로만 비교된다).
엑셀의 '개설대학 및 개설학과(부)' 끝에 붙은 한 글자(예: '전자공학부 A')는 학부 안 트랙 표시라
학과 이름에서 떼어 track 으로 둔다 (개설 강좌의 학과 이름에는 없다).
"""
import json
import re
import shutil
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent / "data"
OUT = ROOT / "rules" / "relative_grading"
HEADER = "교과목 번호"


def _sheet(path):
    return next(ws for ws in openpyxl.load_workbook(path, read_only=True).worksheets
                if any(HEADER in [str(v).strip() for v in r if v] for r in ws.iter_rows(max_row=10, values_only=True)))


def read_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows) if HEADER in [str(v).strip() for v in r if v])
    cols = {str(v).strip(): i for i, v in enumerate(rows[start]) if v}
    for r in rows[start + 1:]:
        code = (r[cols[HEADER]] or "").strip()
        if re.fullmatch(r"[A-Z]{4}\d{4}", code):
            yield {k: r[i] for k, i in cols.items()}


def detect(ws):
    """제목 '2026.2학기 …'·'※ 작성 기준일: 2026. 7. 10.(금)' → ('2026-2', '2026-07-10')"""
    text = " ".join(str(v) for r in ws.iter_rows(max_row=8, values_only=True) for v in r if v)
    term = re.search(r"(20\d\d)\s*[.년-]\s*([12])\s*학기", text)
    date = re.search(r"기준일\s*:?\s*(20\d\d)\.\s*(\d{1,2})\.\s*(\d{1,2})", text)
    return (f"{term[1]}-{term[2]}" if term else None,
            f"{date[1]}-{int(date[2]):02d}-{int(date[3]):02d}" if date else None)


def split_dept(text):
    """'IT대학 전자공학부 A' → ('IT대학', '전자공학부', 'A')"""
    college, _, dept = re.sub(r"\s+", " ", text.strip()).partition(" ")
    m = re.fullmatch(r"(.+) ([A-Z])", dept)
    return (college, m[1], m[2]) if m else (college, dept, None)


def main(xlsx, term=None, notice_date=None):
    ws = _sheet(xlsx)
    found_term, found_date = detect(ws)
    term, notice_date = term or found_term, notice_date or found_date
    if not (term and re.fullmatch(r"20\d\d-[12]", term)) or not notice_date:
        sys.exit(f"학기·작성 기준일을 엑셀에서 못 읽었어요 (학기={term}, 기준일={notice_date}) — 직접 넣어 주세요.\n" + __doc__)

    tt = ROOT / "timetable" / f"{term}.json"
    sections = json.loads(tt.read_text(encoding="utf-8"))["courses"] if tt.exists() else []
    if not sections:
        print(f"※ {tt.name} 없음 — 분반 대조 없이 학과 이름으로만 비교 (scripts/fetch_timetable.py 로 먼저 받으면 더 정확)")
    courses = []
    for r in read_rows(ws):
        college, dept, track = split_dept(r["개설대학 및 개설학과(부)"])
        code = r[HEADER].strip()
        matched = [s for s in sections if s["code"] == code and s["college"] == college and s["dept"] == dept]
        courses.append({
            "code": code, "name": str(r["교과목명"]).strip(), "college": college, "dept": dept, "track": track,
            "grade": int(r["개설학년"]) if r["개설학년"] else None, "credits": r["학점"], "type": r["교과구분"],
            "dept_codes": sorted({s["dept_code"] for s in matched}),
            "sections": sorted(s["section"] for s in matched),
        })

    # 원본도 학기 이름으로 보관 (다시 변환하거나 대조할 때)
    source = ROOT / "sources" / f"knu_{term}_relative_grading.xlsx"
    if Path(xlsx).resolve() != source.resolve():
        shutil.copyfile(xlsx, source)
    OUT.mkdir(parents=True, exist_ok=True)
    out = OUT / f"{term}.json"
    out.write_text(json.dumps({
        "term": term,
        "notice_date": notice_date,
        "source": source.name,
        "courses": courses,
    }, ensure_ascii=False, indent=1), encoding="utf-8")
    unmatched = sum(not c["sections"] for c in courses)
    print(f"{term} (기준일 {notice_date}): {len(courses)}과목 → {out.relative_to(ROOT.parent)} "
          f"(개설 강좌에서 분반을 못 찾은 지정 {unmatched}개 — 이름으로만 비교)")


if __name__ == "__main__":
    if len(sys.argv) not in (2, 4):
        sys.exit(__doc__)
    main(*sys.argv[1:])
